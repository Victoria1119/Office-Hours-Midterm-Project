# make_docs.py
# Reads schedule_data.txt (made by the C program) and creates
# schedule.docx (Word) and schedule.xlsx (Excel).
# The Makefile runs this after the C program.
#
# Needs: pip install --user python-docx openpyxl

import sys

try:
    from docx import Document
    from openpyxl import Workbook
except ImportError:
    print("Missing libraries. Run: pip install --user python-docx openpyxl")
    sys.exit(1)

# each line: professor|department|CLASS or OFFICE|course or days|times|location
rows = []
with open("schedule_data.txt") as f:
    for line in f:
        rows.append(line.strip().split("|"))

# list of professors in the order they appear
professors = []
for r in rows:
    if r[0] not in professors:
        professors.append(r[0])

# ---------- Word: one section per professor ----------
doc = Document()
for name in professors:
    their_rows = [r for r in rows if r[0] == name]
    doc.add_heading(name + " - " + their_rows[0][1], level=1)

    for kind, title, first_col in [("CLASS", "CLASS SCHEDULE", "Course"),
                                   ("OFFICE", "OFFICE HOURS", "Days")]:
        doc.add_heading(title, level=2)
        table = doc.add_table(rows=1, cols=3)
        table.style = "Table Grid"
        table.rows[0].cells[0].text = first_col
        table.rows[0].cells[1].text = "Times"
        table.rows[0].cells[2].text = "Location"
        for r in their_rows:
            if r[2] == kind:
                cells = table.add_row().cells
                cells[0].text = r[3]
                cells[1].text = r[4]
                cells[2].text = r[5]

    doc.add_paragraph("NOTE: Instructors are not to be interrupted during class times.")
doc.save("schedule.docx")

# ---------- Excel: one sheet with every row ----------
wb = Workbook()
ws = wb.active
ws.title = "Schedule"
ws.append(["Professor", "Department", "Type", "Course / Days", "Times", "Location"])
for r in rows:
    r_type = "Class" if r[2] == "CLASS" else "Office Hours"
    ws.append([r[0], r[1], r_type, r[3], r[4], r[5]])
wb.save("schedule.xlsx")

print("Created schedule.docx and schedule.xlsx")